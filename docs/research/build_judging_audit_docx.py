#!/usr/bin/env python3
"""Build the ContextShield SIH judging-readiness audit as a polished DOCX."""

from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs" / "research" / "report-source.md"
OUTPUT = ROOT / "deliverables" / "ContextShield-Judging-Readiness-Audit-2026-09-03.docx"
HELPERS = Path(
    "/Users/arindam/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.826.12353/skills/documents/scripts"
)
sys.path.insert(0, str(HELPERS))
from table_geometry import apply_table_geometry, column_widths_from_weights  # noqa: E402


NAVY = "17324D"
BLUE = "2E74B5"
LIGHT_BLUE = "EAF2F8"
TEAL = "087F5B"
LIGHT_TEAL = "E7F4EF"
AMBER = "A45A00"
LIGHT_AMBER = "FFF3D6"
RED = "A63D2F"
LIGHT_RED = "FDEDEA"
INK = "24313D"
MUTED = "5F6F7C"
LINE = "C9D4DE"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{side}"
        edge = tc_borders.find(qn(tag))
        if edge is None:
            edge = OxmlElement(tag)
            tc_borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK, size: float = 8.4) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    add_inline_runs(p, text, default_size=size, default_color=color, default_bold=bold)


def add_hyperlink(paragraph, text: str, url: str, *, size: float = 10.5) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\(https?://[^)]+\))")


def add_inline_runs(
    paragraph,
    text: str,
    *,
    default_size: float = 10.5,
    default_color: str = INK,
    default_bold: bool = False,
) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            run.bold = default_bold
            run.font.size = Pt(default_size)
            run.font.color.rgb = RGBColor.from_string(default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.size = Pt(default_size)
            run.font.color.rgb = RGBColor.from_string(default_color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Aptos Mono"
            run.font.size = Pt(default_size - 0.5)
            run.font.color.rgb = RGBColor.from_string(NAVY)
            run._element.get_or_add_rPr().append(_run_shading("EDF1F4"))
        else:
            link = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2), size=default_size)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = default_bold
        run.font.size = Pt(default_size)
        run.font.color.rgb = RGBColor.from_string(default_color)


def _run_shading(fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    return shd


def add_rule(paragraph, color: str = BLUE, size: int = 18) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_page_number(paragraph) -> None:
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run("\t")
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Title", 27, NAVY, 0, 8),
        ("Subtitle", 12.5, MUTED, 0, 12),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Audit Label" not in styles:
        label = styles.add_style("Audit Label", WD_STYLE_TYPE.PARAGRAPH)
    else:
        label = styles["Audit Label"]
    label.font.name = "Calibri"
    label.font.size = Pt(8.5)
    label.font.bold = True
    label.font.color.rgb = RGBColor.from_string(TEAL)
    label.paragraph_format.space_after = Pt(3)
    label.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.05


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def build_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    for header in (section.header,):
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("CONTEXTSHIELD  |  SIH 2026")
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(NAVY)
        r2 = p.add_run("                                             JUDGING-READINESS AUDIT")
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = RGBColor.from_string(MUTED)
        add_rule(p, LINE, 6)

    for footer in (section.footer,):
        p = footer.paragraphs[0]
        p.add_run("Evidence audit • 3 September 2026")
        p.runs[0].font.size = Pt(8.5)
        p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
        add_page_number(p)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph(style="Audit Label")
    p.add_run("INTERNAL DECISION BRIEF  •  EVIDENCE-LED  •  JUDGE-SAFE")

    p = doc.add_paragraph(style="Title")
    p.add_run("ContextShield\njudging-readiness audit")

    p = doc.add_paragraph(style="Subtitle")
    p.add_run(
        "Where the SIH 2026 solution stands today — what is proved, what is controlled, "
        "and what will move the score."
    )
    add_rule(p, BLUE, 20)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Assessment date  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run("3 September 2026\n")
    r = p.add_run("Problem statement  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run("SIH26171 — On-device Visual Perception for Light-weight Browser Agents\n")
    r = p.add_run("Evidence base  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.add_run("Repository inspection, fresh test run, current primary research, official product guidance")

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    widths = column_widths_from_weights([1.1, 2.3, 2.6], 9360)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=9360,
        indent_dxa=160,
        cell_margins_dxa={"top": 150, "bottom": 150, "start": 160, "end": 160},
    )
    cells = table.rows[0].cells
    set_cell_shading(cells[0], NAVY)
    set_cell_text(cells[0], "71/100", bold=True, color=WHITE, size=18)
    set_cell_shading(cells[1], LIGHT_TEAL)
    set_cell_text(cells[1], "Competition-worthy core", bold=True, color=TEAL, size=10.5)
    set_cell_shading(cells[2], LIGHT_AMBER)
    set_cell_text(cells[2], "Not yet independently validated or deployment-ready", bold=True, color=AMBER, size=9.4)
    for c in cells:
        set_cell_border(c, WHITE, 4)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Bottom line")
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_inline_runs(
        p,
        "The architecture is the strongest part. The largest score risk is the gap between "
        "controlled 100% regression results and independent, pixel-level, real-world evidence.",
        default_size=11,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Recommended positioning")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.22)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_together = True
    r = p.add_run(
        "“A working privacy-boundary architecture with controlled proof and a clear path to independent validation.”"
    )
    r.italic = True
    r.font.size = Pt(11.5)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    doc.add_page_break()


def style_table(table, header: bool = True, font_size: float = 8.4) -> None:
    for row_idx, row in enumerate(table.rows):
        prevent_row_split(row)
        if row_idx == 0 and header:
            set_repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell, LINE, 5)
            if row_idx == 0 and header:
                set_cell_shading(cell, NAVY)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string(WHITE)
                        run.font.size = Pt(font_size)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F5F8FA")


def add_markdown_table(doc: Document, lines: list[str], table_index: int) -> None:
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c) for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    col_count = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            set_cell_text(
                table.cell(i, j),
                value,
                bold=i == 0 or (i == len(rows) - 1 and value.startswith("**")),
                color=WHITE if i == 0 else INK,
                size=8.0 if col_count >= 4 else 8.4,
            )
    if col_count == 3 and table_index == 1:
        weights = [1.85, 0.8, 4.0]
    elif col_count == 4:
        weights = [1.45, 1.5, 1.8, 2.25]
    elif col_count == 3:
        weights = [1.5, 2.0, 3.4]
    else:
        weights = [1] * col_count
    widths = column_widths_from_weights(weights, 9360)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=9360,
        indent_dxa=120,
        cell_margins_dxa={"top": 80, "bottom": 80, "start": 120, "end": 120},
    )
    style_table(table, font_size=8.0 if col_count >= 4 else 8.4)
    if table_index == 1:
        # Give the total row a clear visual close.
        for cell in table.rows[-1].cells:
            set_cell_shading(cell, LIGHT_BLUE)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = False
    add_inline_runs(p, text, default_size=10.5)


def create_numbering(doc: Document, start: int = 1) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = (max(abstract_ids) + 1) if abstract_ids else 1
    num_id = (max(num_ids) + 1) if num_ids else 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start_el = OxmlElement("w:start")
    start_el.set(qn("w:val"), str(start))
    lvl.append(start_el)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "400")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "400")
    ind.set(qn("w:hanging"), "260")
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc: Document, text: str, ordered: bool, num_id: int | None = None) -> None:
    p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
    if ordered and num_id is not None:
        p_pr = p._p.get_or_add_pPr()
        num_pr = p_pr.get_or_add_numPr()
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_pr.append(ilvl)
        num_id_el = OxmlElement("w:numId")
        num_id_el.set(qn("w:val"), str(num_id))
        num_pr.append(num_id_el)
    p.paragraph_format.keep_together = False
    add_inline_runs(p, text, default_size=10.3)


def add_status_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "7")
    left.set(qn("w:color"), TEAL)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    add_inline_runs(p, text, default_size=10.2, default_color=TEAL, default_bold=True)


def render_markdown(doc: Document, content: str) -> None:
    lines = content.splitlines()
    i = 0
    table_index = 0
    ordered_num_id: int | None = None
    last_order_number: int | None = None
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            add_body_paragraph(doc, " ".join(x.strip() for x in paragraph_buffer))
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            flush_paragraph()
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("**Audience:") or line.startswith("**Assessment date:") or line.startswith("**Problem:") or line.startswith("**Decision:"):
            i += 1
            continue
        if line.startswith("|"):
            flush_paragraph()
            table_lines = []
            while i < len(lines) and lines[i].rstrip().startswith("|"):
                table_lines.append(lines[i].rstrip())
                i += 1
            table_index += 1
            add_markdown_table(doc, table_lines, table_index)
            continue
        heading = re.match(r"^(#{2,3})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            text = heading.group(2)
            if text in {
                "Headline scorecard",
                "Architecturally strong and directly inspectable",
                "Rubric-by-rubric assessment",
                "Priority plan before judging",
                "Claim-to-source ledger",
            }:
                doc.add_page_break()
            if heading.group(1) == "##":
                p = doc.add_paragraph(style="Heading 1")
            else:
                p = doc.add_paragraph(style="Heading 2")
            add_inline_runs(p, text, default_size=16 if heading.group(1) == "##" else 13, default_color=BLUE, default_bold=True)
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if numbered:
            flush_paragraph()
            order_number = int(numbered.group(1))
            if ordered_num_id is None or last_order_number is None or order_number != last_order_number + 1:
                ordered_num_id = create_numbering(doc, start=order_number)
            add_list_item(doc, numbered.group(2), True, ordered_num_id)
            last_order_number = order_number
            i += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            flush_paragraph()
            add_list_item(doc, bullet.group(1), False)
            i += 1
            continue
        if line.startswith("**Judge-safe status:**"):
            flush_paragraph()
            add_status_paragraph(doc, line.replace("**Judge-safe status:**", "Judge-safe status —"))
            i += 1
            continue
        paragraph_buffer.append(line)
        i += 1
    flush_paragraph()


def add_disclaimer(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    add_rule(p, LINE, 6)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(
        "Internal analytical estimate. This is not an official SIH score, legal opinion, or independent certification."
    )
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(MUTED)


def normalize_even_page_references(path: Path) -> None:
    """Make LibreOffice and Word reuse the same header/footer part on even pages."""
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    import xml.etree.ElementTree as ET

    with ZipFile(path, "r") as source_zip:
        members = {name: source_zip.read(name) for name in source_zip.namelist()}
    root = ET.fromstring(members["word/document.xml"])
    for sect_pr in root.findall(f".//{{{w_ns}}}sectPr"):
        for tag in ("headerReference", "footerReference"):
            refs = sect_pr.findall(f"{{{w_ns}}}{tag}")
            default_ref = next((x for x in refs if x.get(f"{{{w_ns}}}type") == "default"), None)
            even_ref = next((x for x in refs if x.get(f"{{{w_ns}}}type") == "even"), None)
            if default_ref is not None and even_ref is not None:
                even_ref.set(f"{{{r_ns}}}id", default_ref.get(f"{{{r_ns}}}id"))
    members["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=path.parent) as tmp:
        temp_path = Path(tmp.name)
    try:
        with ZipFile(temp_path, "w", ZIP_DEFLATED) as target_zip:
            for name, data in members.items():
                target_zip.writestr(name, data)
        temp_path.replace(path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def main() -> None:
    content = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    configure_section(doc.sections[0])
    configure_styles(doc)
    build_header_footer(doc)
    add_title_page(doc)
    render_markdown(doc, content)
    add_disclaimer(doc)

    doc.core_properties.title = "ContextShield judging-readiness audit"
    doc.core_properties.subject = "SIH 2026 rubric assessment and evidence audit"
    doc.core_properties.author = "ContextShield team"
    doc.core_properties.keywords = "ContextShield, SIH 2026, judging rubric, privacy, browser agent"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
